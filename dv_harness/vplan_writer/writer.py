"""vplan_writer/writer.py -- writes a real, openable .xlsx vPlan workbook
from a protocol-agnostic list of structured verification-item dicts.

Design source: this round's DESIGN SPEC (see the task/session record), which
itself synthesizes:
  - .claude/skills/CORE/vplan-core/SKILL.md (generic vPlan authority)
  - .claude/skills/USB/usb-vplan/SKILL.md (USB-specific mandatory columns,
    no-spec-book "TBD-spec" constraint, pattern/task provenance discipline)
  - .claude/agents/IP_UVM_DV_Gen.md's "The vPlan" section (four-sheet target,
    the 15-field row layout, colour-coded "covered by", the validation
    rules: every pattern name has a matching file, every task name is a real
    declaration, every pattern is in the run-time dispatcher (all 3
    mandatory), AND every `constraint items` entry names a constraint that
    actually exists in the written SV source (4th rule, added 2026-09-01,
    implemented here as an OPTIONAL evidence-gated check -- see
    ConstraintNotInSVSourceError/sv_constraint_names below for the
    backward-compatibility ruling))
  - dv_harness/uvm_generator/generator.py's typed-error convention
    (`class XxxError(ValueError)` with `.reason`/`.detail`, evidence-required
    fields, additive/backward-compatible schema extensions, fail-fast one
    raise per call)

This module never hardcodes a protocol's directory layout: build_evidence_context
takes paths/regexes as arguments so a different IP (not just USB) can supply
its own pattern_dir/dispatcher_file/task_declaration_sources.
"""

from __future__ import annotations

import glob as _glob_module
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Dict, FrozenSet, List, Optional, Sequence, Tuple
from uuid import uuid4

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

# ---------------------------------------------------------------------------
# Enums (module-level constants -- single source of truth for both the
# schema validator and the coverage-summary sheet builder, never re-typed
# as literal strings in two places).
# ---------------------------------------------------------------------------

COVERED_BY_STATES: Tuple[str, ...] = ("covered", "PARTIAL", "NOT COVERED", "N/A", "DEFERRED")
RANDOM_OR_DIRECTED_STATES: Tuple[str, ...] = ("random", "directed")
BLOCKED_ON_STATES: Tuple[str, ...] = ("INFORMATION", "EFFORT")

# Credit rule (module-level constant, easy to extend): mirrors
# spec_coverage_audit.py's VERIFIED/approved-WAIVED/NOT_APPLICABLE credit
# philosophy applied at item granularity -- PARTIAL/NOT COVERED/DEFERRED
# never credit, an explicit auditable honesty rule instead of an implicit one.
CREDITED_STATES: FrozenSet[str] = frozenset({"covered", "N/A"})

# 15-column layout from .claude/agents/IP_UVM_DV_Gen.md's "The vPlan" section,
# in exact order: (header text, VPlanItem field name).
_VPLAN_COLUMNS: Tuple[Tuple[str, str], ...] = (
    ("ID", "req_id"),
    ("Feature area", "feature_area"),
    ("Verification item", "verification_item"),
    ("testing pattern name", "pattern_name"),
    ("command.txt task name", "task_name"),
    ("suite", "suite"),
    ("covered by", "covered_by"),
    ("testing pattern description", "description"),
    ("spec section", "spec_section"),
    ("constraint items", "constraint_items"),
    ("random or directed", "random_or_directed"),
    ("mode/speed", "mode_speed"),
    ("instance", "instance"),
    ("checkers active", "checkers_active"),
    ("notes", "notes"),
)

# Fixed fill map for the "covered by" column (col G), per the agent file's
# colour-coded convention.
_COVERED_BY_FILL: Dict[str, str] = {
    "covered": "C6EFCE",       # green
    "PARTIAL": "FFEB9C",       # amber
    "NOT COVERED": "FFC7CE",   # red
    "N/A": "D9D9D9",           # grey
    "DEFERRED": "BDD7EE",      # blue
}

_BANNER_FILL_COLOR = "DDEBF7"

# Required keys every VPlanItem dict must carry (as a key -- value may
# legitimately be None for pattern_name/task_name/blocked_on/blocked_reason
# per the conditional rules enforced separately below).
_REQUIRED_KEYS: Tuple[str, ...] = (
    "req_id", "feature_area", "verification_item", "pattern_name", "task_name",
    "suite", "covered_by", "description", "spec_section", "constraint_items",
    "random_or_directed", "mode_speed", "instance", "checkers_active",
    "blocked_on", "blocked_reason",
)

# Fields that must be present AND a non-empty string (independent of
# covered_by state).
_REQUIRED_NON_EMPTY_STR_FIELDS: Tuple[str, ...] = (
    "req_id", "feature_area", "verification_item", "suite", "description",
    "spec_section", "mode_speed", "instance",
)


# ---------------------------------------------------------------------------
# Evidence context
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class VPlanEvidenceContext:
    """The three ground-truth sets the three mandatory validation rules from
    .claude/agents/IP_UVM_DV_Gen.md check against, plus one OPTIONAL fourth
    set (sv_constraint_names) for the 4th rule added to that same doc section
    on 2026-09-01 ("every `constraint items` entry names a constraint that
    actually exists in the written SV source"). The provenance fields below
    (pattern_dir/pattern_glob/dispatcher_file/task_declaration_sources/
    constraint_declaration_sources) are additive metadata carried only so a
    raised typed error's `.detail` can cite exactly where a ground-truth set
    came from -- they do not affect validation logic, which reads only the
    frozenset fields.

    sv_constraint_names is None (not an empty frozenset) when the caller
    supplied no constraint-declaration evidence at all -- RULING
    (vplan-4th-rule-implementation, 2026-09-01): the 4th rule is implemented
    as an OPTIONAL, evidence-gated check, exactly mirroring the existing
    known_check_names/UnknownCheckerNameError precedent, not an
    unconditionally mandatory one. Making it unconditional would break
    additive/backward-compatibility for every existing caller (CLI usage,
    STAGE_GATES['VPLAN'] JSON payloads, existing tests) that does not yet
    supply SV constraint-declaration evidence -- None means "rule not
    requested, skip"; a non-None (possibly built from a real scan) frozenset
    means "rule requested, enforce strictly, fail-closed on any mismatch",
    same fail-closed discipline as the three mandatory rules.
    """
    pattern_files: FrozenSet[str]        # stems of real files found under pattern_dir
    dispatcher_patterns: FrozenSet[str]  # pattern names extracted from the runtime dispatcher file
    declared_tasks: FrozenSet[str]       # task/test/vseq names extracted from task-declaration sources
    sv_constraint_names: Optional[FrozenSet[str]] = None  # `constraint <name>` decls found in SV source; None = rule skipped

    pattern_dir: str = ""
    pattern_glob: str = ""
    dispatcher_file: str = ""
    task_declaration_sources: Tuple[str, ...] = ()
    constraint_declaration_sources: Tuple[str, ...] = ()


class EvidenceSourceEmptyError(ValueError):
    """Raised by build_evidence_context (and re-checked defensively by
    write_vplan_workbook before validating any item) when pattern_files,
    dispatcher_patterns, or declared_tasks comes back empty. Refuses to run
    per-item validation against a suspiciously-empty ground-truth set --
    almost certainly a wrong path/glob/regex, not real environment state.
    detail: {"which": "pattern_files"|"dispatcher_patterns"|"declared_tasks",
    "source_path_or_paths": str|list[str]}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


def build_evidence_context(
    *,
    pattern_dir: "str | Path",
    pattern_glob: str = "*.txt",
    dispatcher_file: "str | Path",
    dispatcher_pattern_regex: str = r'"(?P<pattern>[A-Za-z0-9_]+)"',
    task_declaration_sources: Optional[List["str | Path"]] = None,
    task_declaration_regex: str = r'\btask\s+automatic\s+(?P<task>[A-Za-z0-9_]+)\b|\bclass\s+(?P<task2>[A-Za-z0-9_]+)\s+extends\b',
    known_task_names: Optional[FrozenSet[str]] = None,
    constraint_declaration_sources: Optional[List["str | Path"]] = None,
    constraint_declaration_regex: str = r'\bconstraint\s+(?P<constraint>[A-Za-z0-9_]+)\b',
    known_constraint_names: Optional[FrozenSet[str]] = None,
) -> VPlanEvidenceContext:
    """Scans real files (glob + regex text scan -- same evidence-by-citation
    discipline as generator.py's evidence-string requirement, not a full SV
    parser) to build the three ground-truth sets the three mandatory
    validation rules check against. Raises EvidenceSourceEmptyError if any
    resulting set is empty -- an empty ground-truth set almost always means
    a wrong path/regex was wired, not that zero real patterns/tasks exist,
    and letting validation proceed would blame every item's mismatch on the
    item instead of the environment.

    Protocol-agnostic: USB supplies pattern_dir=.../uvm/patterns,
    dispatcher_file=.../dv_uvm_pattern_pool.svh,
    task_declaration_sources=[.../tb/tests/*.sv]; a different IP supplies
    its own paths/regexes.

    constraint_declaration_sources/constraint_declaration_regex/
    known_constraint_names build the OPTIONAL 4th ground-truth set
    (sv_constraint_names) for the 2026-09-01 "constraint items exist in SV
    source" rule -- mirrors task_declaration_sources/known_task_names'
    scan-vs-bypass shape exactly. Neither supplied -> sv_constraint_names
    stays None and ConstraintNotInSVSourceError's check is skipped entirely
    (same optional-check precedent as known_check_names below). Either
    supplied but the resulting set comes back empty -> EvidenceSourceEmptyError
    (which="sv_constraint_names"), same fail-closed discipline as the three
    mandatory sets: a caller that explicitly asked for this check getting
    zero constraints back almost always means a wrong path/regex, not that
    the SV source truly declares none.
    """
    pattern_dir_path = Path(pattern_dir)
    pattern_files = frozenset(
        p.stem for p in pattern_dir_path.glob(pattern_glob)
    ) if pattern_dir_path.is_dir() else frozenset()
    if not pattern_files:
        raise EvidenceSourceEmptyError("EVIDENCE_SOURCE_EMPTY", {
            "which": "pattern_files", "source_path_or_paths": str(pattern_dir_path),
        })

    dispatcher_path = Path(dispatcher_file)
    dispatcher_text = (
        dispatcher_path.read_text(encoding="utf-8", errors="replace")
        if dispatcher_path.is_file() else ""
    )
    dispatcher_patterns = frozenset(
        m.group("pattern") for m in re.finditer(dispatcher_pattern_regex, dispatcher_text)
    )
    if not dispatcher_patterns:
        raise EvidenceSourceEmptyError("EVIDENCE_SOURCE_EMPTY", {
            "which": "dispatcher_patterns", "source_path_or_paths": str(dispatcher_path),
        })

    if known_task_names is not None:
        declared_tasks = frozenset(known_task_names)
        scanned_sources: List[str] = []
    else:
        raw_sources = task_declaration_sources or []
        scanned_sources = []
        for src in raw_sources:
            scanned_sources.extend(sorted(_glob_module.glob(str(src))))
        declared: set = set()
        for path_str in scanned_sources:
            text = Path(path_str).read_text(encoding="utf-8", errors="replace")
            for m in re.finditer(task_declaration_regex, text):
                name = m.group("task") if "task" in m.groupdict() and m.group("task") else None
                if name is None and "task2" in m.groupdict():
                    name = m.group("task2")
                if name:
                    declared.add(name)
        declared_tasks = frozenset(declared)

    if not declared_tasks:
        raise EvidenceSourceEmptyError("EVIDENCE_SOURCE_EMPTY", {
            "which": "declared_tasks",
            "source_path_or_paths": scanned_sources or [str(s) for s in (task_declaration_sources or [])],
        })

    # 4th (OPTIONAL) ground-truth set: `constraint <name>` declarations found
    # in real SV source, for the 2026-09-01 "constraint items exist in SV
    # source" rule. None (not an empty frozenset) when the caller supplied
    # neither known_constraint_names nor constraint_declaration_sources --
    # that means "rule not requested", handled by validate_items skipping
    # the check entirely, never by silently treating an unrequested check as
    # a vacuously-passing empty set.
    constraint_scanned_sources: List[str] = []
    if known_constraint_names is not None:
        sv_constraint_names: Optional[FrozenSet[str]] = frozenset(known_constraint_names)
    elif constraint_declaration_sources is not None:
        for src in constraint_declaration_sources:
            constraint_scanned_sources.extend(sorted(_glob_module.glob(str(src))))
        found_constraints: set = set()
        for path_str in constraint_scanned_sources:
            text = Path(path_str).read_text(encoding="utf-8", errors="replace")
            for m in re.finditer(constraint_declaration_regex, text):
                name = m.group("constraint")
                if name:
                    found_constraints.add(name)
        sv_constraint_names = frozenset(found_constraints)
    else:
        sv_constraint_names = None

    if sv_constraint_names is not None and not sv_constraint_names:
        raise EvidenceSourceEmptyError("EVIDENCE_SOURCE_EMPTY", {
            "which": "sv_constraint_names",
            "source_path_or_paths": constraint_scanned_sources or [str(s) for s in (constraint_declaration_sources or [])],
        })

    return VPlanEvidenceContext(
        pattern_files=pattern_files,
        dispatcher_patterns=dispatcher_patterns,
        declared_tasks=declared_tasks,
        sv_constraint_names=sv_constraint_names,
        pattern_dir=str(pattern_dir_path),
        pattern_glob=pattern_glob,
        dispatcher_file=str(dispatcher_path),
        task_declaration_sources=tuple(scanned_sources),
        constraint_declaration_sources=tuple(constraint_scanned_sources),
    )


# ---------------------------------------------------------------------------
# Typed errors (exact .reason/.detail convention from generator.py)
# ---------------------------------------------------------------------------

class VPlanSchemaError(ValueError):
    """Raised by validate_items when a VPlanItem dict is malformed -- shape,
    missing required field, or invalid enum value -- BEFORE any evidence
    check runs. One class, many reason codes (same convention as
    generator.py's MalformedScoreboardCheckError, which is keyed off a
    single top-level manifest key with many distinct malformed-shape reason
    codes). reason in:
      MISSING_REQUIRED_FIELD | DUPLICATE_REQ_ID | INVALID_COVERED_BY_STATE |
      INVALID_RANDOM_OR_DIRECTED | MISSING_PATTERN_OR_TASK_FOR_COVERED_STATE |
      MISSING_BLOCKED_CLASSIFICATION | MISSING_CONSTRAINT_EVIDENCE_FOR_RANDOM_PATTERN
    detail: {"req_id": <id or None>, "row_index": int, "field": str, "value": Any}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


class UnresolvedPatternFileError(ValueError):
    """Raised by validate_items when an item's pattern_name has no matching
    file under evidence.pattern_files -- the agent file's rule 1 ("every
    pattern name has a matching file"). Never a silently-accepted guessed
    pattern name (cf. command_inventory.csv's real COMMAND/SOURCE columns as
    the only legitimate ground truth for a USB pattern name).
    detail: {"req_id": str, "pattern_name": str, "pattern_dir": str, "pattern_glob": str}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


class UnknownTaskDeclarationError(ValueError):
    """Raised by validate_items when an item's task_name is not present in
    evidence.declared_tasks -- rule 2 ("every task name is a real
    declaration"). detail: {"req_id": str, "task_name": str,
    "scanned_sources": list[str]}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


class PatternNotInDispatcherError(ValueError):
    """Raised by validate_items when an item's pattern_name is not present
    in evidence.dispatcher_patterns -- rule 3 ("every pattern is in the
    run-time dispatcher"). A pattern file existing on disk with no
    dispatcher entry cannot actually be selected at run time (+PATTERN),
    so listing it as covered/PARTIAL would be false. detail: {"req_id": str,
    "pattern_name": str, "dispatcher_file": str}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


class UnknownCheckerNameError(ValueError):
    """Raised by validate_items ONLY when the caller supplied
    known_check_names (typically generator.py scoreboard_rules[].check_name
    values for a manifest-generated environment) and an item's
    checkers_active entry is not in that set. Optional check: legacy
    pre-generator BFM environments have no such manifest, so a None
    known_check_names skips this rule entirely rather than failing closed.
    detail: {"req_id": str, "checker_name": str, "known_check_names": list[str]}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


class ConstraintNotInSVSourceError(ValueError):
    """Raised by validate_items when an item's constraint_items entry does
    not name a constraint that actually exists in the written SV source --
    the 4th vPlan validation rule added to .claude/agents/IP_UVM_DV_Gen.md's
    "The vPlan" section on 2026-09-01: "every `constraint items` entry names
    a constraint that actually exists in the written SV source". Only
    checked when evidence.sv_constraint_names is not None, i.e. the caller
    supplied constraint_declaration_sources or known_constraint_names to
    build_evidence_context (same OPTIONAL-check precedent as
    UnknownCheckerNameError/known_check_names -- a caller with no
    constraint-declaration evidence at all, e.g. a legacy pre-generator BFM
    environment, skips this rule entirely rather than failing closed on
    evidence it never supplied). Once requested, a caller never accepts a
    guessed/aspirational constraint name -- refuses to write rather than
    letting the vPlan's `constraint items` column drift from real testbench
    code.
    detail: {"req_id": str, "constraint_name": str,
    "constraint_declaration_sources": list[str]}"""
    def __init__(self, reason: str, detail: dict):
        super().__init__(reason)
        self.reason = reason
        self.detail = detail


# ---------------------------------------------------------------------------
# Schema validation (pure, no evidence lookups)
# ---------------------------------------------------------------------------

def _validate_item_schema(item: Any, idx: int, seen_req_ids: set) -> None:
    if not isinstance(item, dict):
        raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
            "req_id": None, "row_index": idx, "field": "<row>", "value": item,
        })

    for field_name in _REQUIRED_KEYS:
        if field_name not in item:
            raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
                "req_id": item.get("req_id"), "row_index": idx, "field": field_name, "value": None,
            })

    req_id = item["req_id"]
    for field_name in _REQUIRED_NON_EMPTY_STR_FIELDS:
        value = item.get(field_name)
        if not isinstance(value, str) or not value.strip():
            raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
                "req_id": req_id, "row_index": idx, "field": field_name, "value": value,
            })

    if req_id in seen_req_ids:
        raise VPlanSchemaError("DUPLICATE_REQ_ID", {
            "req_id": req_id, "row_index": idx, "field": "req_id", "value": req_id,
        })
    seen_req_ids.add(req_id)

    covered_by = item["covered_by"]
    if covered_by not in COVERED_BY_STATES:
        raise VPlanSchemaError("INVALID_COVERED_BY_STATE", {
            "req_id": req_id, "row_index": idx, "field": "covered_by", "value": covered_by,
        })

    random_or_directed = item["random_or_directed"]
    if random_or_directed not in RANDOM_OR_DIRECTED_STATES:
        raise VPlanSchemaError("INVALID_RANDOM_OR_DIRECTED", {
            "req_id": req_id, "row_index": idx, "field": "random_or_directed", "value": random_or_directed,
        })

    pattern_name = item["pattern_name"]
    task_name = item["task_name"]
    if covered_by not in ("N/A", "DEFERRED") and (pattern_name is None or task_name is None):
        raise VPlanSchemaError("MISSING_PATTERN_OR_TASK_FOR_COVERED_STATE", {
            "req_id": req_id, "row_index": idx, "field": "pattern_name/task_name",
            "value": {"pattern_name": pattern_name, "task_name": task_name},
        })

    blocked_on = item["blocked_on"]
    blocked_reason = item["blocked_reason"]
    if covered_by == "NOT COVERED":
        if blocked_on not in BLOCKED_ON_STATES or not blocked_reason or not str(blocked_reason).strip():
            raise VPlanSchemaError("MISSING_BLOCKED_CLASSIFICATION", {
                "req_id": req_id, "row_index": idx, "field": "blocked_on/blocked_reason",
                "value": {"blocked_on": blocked_on, "blocked_reason": blocked_reason},
            })

    constraint_items = item["constraint_items"]
    if not isinstance(constraint_items, list):
        raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
            "req_id": req_id, "row_index": idx, "field": "constraint_items", "value": constraint_items,
        })
    if random_or_directed == "random" and not constraint_items:
        raise VPlanSchemaError("MISSING_CONSTRAINT_EVIDENCE_FOR_RANDOM_PATTERN", {
            "req_id": req_id, "row_index": idx, "field": "constraint_items", "value": constraint_items,
        })

    checkers_active = item["checkers_active"]
    if not isinstance(checkers_active, list):
        raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
            "req_id": req_id, "row_index": idx, "field": "checkers_active", "value": checkers_active,
        })
    if not checkers_active and covered_by not in ("NOT COVERED", "N/A", "DEFERRED"):
        raise VPlanSchemaError("MISSING_REQUIRED_FIELD", {
            "req_id": req_id, "row_index": idx, "field": "checkers_active", "value": checkers_active,
        })


def _validate_item_evidence(
    item: dict, evidence: VPlanEvidenceContext, known_check_names: Optional[FrozenSet[str]],
) -> None:
    req_id = item["req_id"]
    pattern_name = item["pattern_name"]
    task_name = item["task_name"]
    checkers_active = item["checkers_active"]
    constraint_items = item["constraint_items"]

    if pattern_name is not None:
        if pattern_name not in evidence.pattern_files:
            raise UnresolvedPatternFileError("UNRESOLVED_PATTERN_FILE", {
                "req_id": req_id, "pattern_name": pattern_name,
                "pattern_dir": evidence.pattern_dir, "pattern_glob": evidence.pattern_glob,
            })
        if pattern_name not in evidence.dispatcher_patterns:
            raise PatternNotInDispatcherError("PATTERN_NOT_IN_DISPATCHER", {
                "req_id": req_id, "pattern_name": pattern_name,
                "dispatcher_file": evidence.dispatcher_file,
            })

    if task_name is not None:
        if task_name not in evidence.declared_tasks:
            raise UnknownTaskDeclarationError("UNKNOWN_TASK_DECLARATION", {
                "req_id": req_id, "task_name": task_name,
                "scanned_sources": list(evidence.task_declaration_sources),
            })

    if known_check_names is not None:
        for checker_name in checkers_active:
            if checker_name not in known_check_names:
                raise UnknownCheckerNameError("UNKNOWN_CHECKER_NAME", {
                    "req_id": req_id, "checker_name": checker_name,
                    "known_check_names": sorted(known_check_names),
                })

    if evidence.sv_constraint_names is not None:
        for constraint_name in constraint_items:
            if constraint_name not in evidence.sv_constraint_names:
                raise ConstraintNotInSVSourceError("CONSTRAINT_NOT_IN_SV_SOURCE", {
                    "req_id": req_id, "constraint_name": constraint_name,
                    "constraint_declaration_sources": list(evidence.constraint_declaration_sources),
                })


def validate_items(
    items: List[dict],
    evidence: VPlanEvidenceContext,
    *,
    known_check_names: Optional[FrozenSet[str]] = None,
) -> None:
    """Raises on the first violation found. Validation order (fail-fast, one
    raise per call -- matches generator.py's convention of raising
    immediately on the first bad manifest entry rather than aggregating):
      (1) EvidenceSourceEmptyError pre-flight on the context itself,
      (2) VPlanSchemaError over every item in list order,
      (3) evidence checks (UnresolvedPatternFileError ->
          PatternNotInDispatcherError -> UnknownTaskDeclarationError ->
          UnknownCheckerNameError -> ConstraintNotInSVSourceError) over every
          item in list order. ConstraintNotInSVSourceError is only raised
          when evidence.sv_constraint_names is not None (see
          build_evidence_context's constraint_declaration_sources/
          known_constraint_names) -- an OPTIONAL check, same precedent as
          UnknownCheckerNameError/known_check_names.
    Returns None on success. Callable standalone by gates.py or an audit
    script without writing a workbook.

    A caller wanting *all* failures at once can wrap this per-item in its
    own try/except loop -- batching is a caller-side concern, not internal
    complexity this module takes on.
    """
    for which, source, source_path in (
        ("pattern_files", evidence.pattern_files, evidence.pattern_dir),
        ("dispatcher_patterns", evidence.dispatcher_patterns, evidence.dispatcher_file),
        ("declared_tasks", evidence.declared_tasks, list(evidence.task_declaration_sources)),
    ):
        if not source:
            raise EvidenceSourceEmptyError("EVIDENCE_SOURCE_EMPTY", {
                "which": which, "source_path_or_paths": source_path,
            })

    seen_req_ids: set = set()
    for idx, item in enumerate(items):
        _validate_item_schema(item, idx, seen_req_ids)

    for item in items:
        _validate_item_evidence(item, evidence, known_check_names)


# ---------------------------------------------------------------------------
# Derived data (feature-area ordering, counts, gap ranking) -- pure helpers
# shared between the workbook writer and its result summary.
# ---------------------------------------------------------------------------

def _feature_area_order(items: List[dict]) -> List[str]:
    """First-seen order of feature_area across items -- preserves the
    caller's intended narrative order rather than alphabetizing."""
    order: List[str] = []
    seen: set = set()
    for item in items:
        fa = item["feature_area"]
        if fa not in seen:
            seen.add(fa)
            order.append(fa)
    return order


def _index_to_letters(i: int) -> str:
    """0-based index -> Excel-style column letters (A, B, ..., Z, AA, AB, ...)
    -- used to letter feature-area sections beyond 26 without collision."""
    i += 1
    letters = ""
    while i > 0:
        i, rem = divmod(i - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _counts_by_state(items: List[dict]) -> Dict[str, int]:
    counts = {s: 0 for s in COVERED_BY_STATES}
    for item in items:
        counts[item["covered_by"]] = counts.get(item["covered_by"], 0) + 1
    return counts


def _coverage_percent(items: List[dict]) -> float:
    if not items:
        return 0.0
    credited = sum(1 for it in items if it["covered_by"] in CREDITED_STATES)
    return 100.0 * credited / len(items)


def _rank_gaps(items: List[dict], feature_order: List[str]) -> List[dict]:
    """Deterministic 3-tier sort (ascending rank = close first): tier 0 =
    PARTIAL, tier 1 = NOT COVERED/blocked on EFFORT, tier 2 = NOT
    COVERED/blocked on INFORMATION -- then by feature_area first-seen
    order, then by req_id."""
    feature_index = {fa: i for i, fa in enumerate(feature_order)}
    candidates: List[Tuple[int, int, str, dict]] = []
    for item in items:
        cb = item["covered_by"]
        if cb == "PARTIAL":
            tier = 0
        elif cb == "NOT COVERED" and item.get("blocked_on") == "EFFORT":
            tier = 1
        elif cb == "NOT COVERED" and item.get("blocked_on") == "INFORMATION":
            tier = 2
        else:
            continue
        candidates.append((tier, feature_index.get(item["feature_area"], 0), item["req_id"], item))
    candidates.sort(key=lambda t: (t[0], t[1], t[2]))

    gaps: List[dict] = []
    for rank, (tier, _, _, item) in enumerate(candidates, start=1):
        if tier == 0:
            why = "PARTIAL — pattern/task exist (%s/%s); close by: %s" % (
                item.get("pattern_name"), item.get("task_name"), item.get("notes") or "",
            )
        elif tier == 1:
            why = "NOT COVERED, blocked on EFFORT — %s" % (item.get("blocked_reason"),)
        else:
            why = "NOT COVERED, blocked on INFORMATION — %s; cannot start until resolved" % (
                item.get("blocked_reason"),
            )
        gaps.append({
            "rank": rank, "req_id": item["req_id"], "feature_area": item["feature_area"],
            "verification_item": item["verification_item"], "covered_by": item["covered_by"],
            "blocked_on": item.get("blocked_on"), "why": why,
        })
    return gaps


def _deferred_rows(items: List[dict], feature_order: List[str]) -> List[dict]:
    """DEFERRED rows listed separately, not mixed into the ranked gaps --
    DEFERRED is a conscious postponement, not an open gap to prioritize now."""
    feature_index = {fa: i for i, fa in enumerate(feature_order)}
    deferred = [it for it in items if it["covered_by"] == "DEFERRED"]
    deferred.sort(key=lambda it: (feature_index.get(it["feature_area"], 0), it["req_id"]))
    return [
        {
            "req_id": it["req_id"], "feature_area": it["feature_area"],
            "verification_item": it["verification_item"], "notes": it.get("notes") or "",
        }
        for it in deferred
    ]


def _notes_with_blocked(item: dict) -> str:
    notes = item.get("notes") or ""
    blocked_on = item.get("blocked_on")
    blocked_reason = item.get("blocked_reason")
    if blocked_on:
        suffix = "blocked_on=%s; blocked_reason=%s" % (blocked_on, blocked_reason)
        return "%s (%s)" % (notes, suffix) if notes else suffix
    return notes


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

@dataclass
class _SheetContext:
    items: List[dict]
    feature_order: List[str]
    feature_letter_map: Dict[str, str]
    counts_by_state: Dict[str, int]
    coverage_percent: float
    gaps_ranked: List[dict]
    deferred_rows: List[dict]
    protocol: str


def _build_verification_plan_sheet(wb: Workbook, ctx: _SheetContext) -> None:
    ws = wb.create_sheet("verification_plan")
    header_font = Font(bold=True)
    n_cols = len(_VPLAN_COLUMNS)

    for col_idx, (header, _field) in enumerate(_VPLAN_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font

    row = 2
    banner_font = Font(bold=True)
    banner_fill = PatternFill(start_color=_BANNER_FILL_COLOR, end_color=_BANNER_FILL_COLOR, fill_type="solid")
    current_feature = None

    for item in ctx.items:
        fa = item["feature_area"]
        if fa != current_feature:
            current_feature = fa
            letter = ctx.feature_letter_map[fa]
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=n_cols)
            banner_cell = ws.cell(row=row, column=1, value="%s. %s" % (letter, fa))
            banner_cell.font = banner_font
            for c in range(1, n_cols + 1):
                ws.cell(row=row, column=c).fill = banner_fill
            row += 1

        for col_idx, (_header, field_name) in enumerate(_VPLAN_COLUMNS, start=1):
            if field_name in ("constraint_items", "checkers_active"):
                value = ", ".join(item.get(field_name) or [])
            elif field_name == "notes":
                value = _notes_with_blocked(item)
            else:
                value = item.get(field_name)
                if value is None:
                    value = ""
            cell = ws.cell(row=row, column=col_idx, value=value)
            if field_name == "covered_by":
                fill_color = _COVERED_BY_FILL.get(item["covered_by"])
                if fill_color:
                    cell.fill = PatternFill(start_color=fill_color, end_color=fill_color, fill_type="solid")
        row += 1

    last_row = row - 1
    if last_row >= 1:
        ws.auto_filter.ref = "A1:%s%d" % (_index_to_letters(n_cols - 1), last_row)
    ws.freeze_panes = "F2"


def _build_coverage_summary_sheet(wb: Workbook, ctx: _SheetContext) -> None:
    ws = wb.create_sheet("coverage_summary")
    bold = Font(bold=True)
    row = 1

    ws.cell(row=row, column=1, value="Table A — Counts by covered-by state").font = bold
    row += 1
    for c, h in enumerate(("State", "Count", "% of total", "Credited toward coverage goal"), start=1):
        ws.cell(row=row, column=c, value=h).font = bold
    row += 1
    total = len(ctx.items)
    for state in COVERED_BY_STATES:
        count = ctx.counts_by_state.get(state, 0)
        pct = (100.0 * count / total) if total else 0.0
        credited = "Yes" if state in CREDITED_STATES else "No"
        ws.cell(row=row, column=1, value=state)
        ws.cell(row=row, column=2, value=count)
        ws.cell(row=row, column=3, value=round(pct, 2))
        ws.cell(row=row, column=4, value=credited)
        row += 1
    ws.cell(row=row, column=1, value="TOTAL").font = bold
    ws.cell(row=row, column=2, value=total).font = bold
    row += 1
    ws.cell(row=row, column=1, value="Coverage %").font = bold
    ws.cell(row=row, column=2, value=round(ctx.coverage_percent, 2)).font = bold
    row += 2

    ws.cell(row=row, column=1, value="Table B — Gaps ranked").font = bold
    row += 1
    for c, h in enumerate(
        ("Rank", "ID", "Feature area", "Verification item", "covered by", "blocked on", "why it is next"), start=1,
    ):
        ws.cell(row=row, column=c, value=h).font = bold
    row += 1
    for gap in ctx.gaps_ranked:
        ws.cell(row=row, column=1, value=gap["rank"])
        ws.cell(row=row, column=2, value=gap["req_id"])
        ws.cell(row=row, column=3, value=gap["feature_area"])
        ws.cell(row=row, column=4, value=gap["verification_item"])
        ws.cell(row=row, column=5, value=gap["covered_by"])
        ws.cell(row=row, column=6, value=gap["blocked_on"] or "")
        ws.cell(row=row, column=7, value=gap["why"])
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Table C — Deferred (visible, not ranked)").font = bold
    row += 1
    for c, h in enumerate(("ID", "Feature area", "Verification item", "notes"), start=1):
        ws.cell(row=row, column=c, value=h).font = bold
    row += 1
    for dr in ctx.deferred_rows:
        ws.cell(row=row, column=1, value=dr["req_id"])
        ws.cell(row=row, column=2, value=dr["feature_area"])
        ws.cell(row=row, column=3, value=dr["verification_item"])
        ws.cell(row=row, column=4, value=dr["notes"])
        row += 1


def _unimplemented_sheet_stub(key: str) -> Callable[[Workbook, _SheetContext], None]:
    def _builder(wb: Workbook, ctx: _SheetContext) -> None:
        raise NotImplementedError(
            "sheet '%s' not implemented in this round -- see dv_harness/vplan_writer follow-up note" % (key,)
        )
    return _builder


# Registry-key dict (not a fixed pair) so "mode_speed_matrix"/"reference"
# can be added later without changing write_vplan_workbook's signature.
# Reserved-but-unimplemented keys are present, mapped to a stub that raises
# NotImplementedError -- a caller asking for them gets a loud, typed
# refusal now, not silent omission later.
_SHEET_BUILDERS: Dict[str, Callable[[Workbook, _SheetContext], None]] = {
    "verification_plan": _build_verification_plan_sheet,
    "coverage_summary": _build_coverage_summary_sheet,
    "mode_speed_matrix": _unimplemented_sheet_stub("mode_speed_matrix"),
    "reference": _unimplemented_sheet_stub("reference"),
}


# ---------------------------------------------------------------------------
# Public write entry point
# ---------------------------------------------------------------------------

@dataclass(frozen=True)
class VPlanWriteResult:
    path: Path
    row_count: int
    coverage_percent: float
    counts_by_state: Dict[str, int]
    gaps_ranked: List[dict]


def write_vplan_workbook(
    items: List[dict],
    *,
    output_path: "str | Path",
    evidence: VPlanEvidenceContext,
    protocol: str,
    known_check_names: Optional[FrozenSet[str]] = None,
    sheets: Tuple[str, ...] = ("verification_plan", "coverage_summary"),
    workbook_title: Optional[str] = None,
) -> VPlanWriteResult:
    """Calls validate_items(items, evidence, known_check_names=...) first --
    refuses to write on any mismatch, raising the exact typed error above,
    never a partially-written file. Only on success does it build the
    workbook (openpyxl) and write it atomically: build to a temp path in
    the same directory, then os.replace() onto output_path, so a crash
    mid-write never leaves a corrupt/half-written .xlsx where a stale-but-
    valid one used to be.

    sheets is a registry-key tuple, not a fixed pair, precisely so
    "mode_speed_matrix"/"reference" can be added later without changing this
    signature. Requesting an unimplemented key raises NotImplementedError
    naming the key -- never silently ignored.
    """
    validate_items(items, evidence, known_check_names=known_check_names)

    for key in sheets:
        if key not in _SHEET_BUILDERS:
            raise NotImplementedError(
                "sheet '%s' is not a registered vplan_writer sheet key (known: %s)"
                % (key, sorted(_SHEET_BUILDERS.keys()))
            )

    feature_order = _feature_area_order(items)
    feature_letter_map = {fa: _index_to_letters(i) for i, fa in enumerate(feature_order)}
    counts_by_state = _counts_by_state(items)
    coverage_percent = _coverage_percent(items)
    gaps_ranked = _rank_gaps(items, feature_order)
    deferred_rows = _deferred_rows(items, feature_order)

    ctx = _SheetContext(
        items=items, feature_order=feature_order, feature_letter_map=feature_letter_map,
        counts_by_state=counts_by_state, coverage_percent=coverage_percent,
        gaps_ranked=gaps_ranked, deferred_rows=deferred_rows, protocol=protocol,
    )

    wb = Workbook()
    wb.remove(wb.active)  # drop the default blank "Sheet"
    if workbook_title:
        wb.properties.title = workbook_title

    for key in sheets:
        _SHEET_BUILDERS[key](wb, ctx)

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = out_path.parent / (".%s.tmp-%s" % (out_path.name, uuid4().hex))
    wb.save(str(tmp_path))
    os.replace(str(tmp_path), str(out_path))

    return VPlanWriteResult(
        path=out_path,
        row_count=len(items),
        coverage_percent=coverage_percent,
        counts_by_state=counts_by_state,
        gaps_ranked=gaps_ranked,
    )
