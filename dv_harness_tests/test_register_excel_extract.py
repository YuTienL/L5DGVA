"""Tests for dv_harness/register_excel_extract.py.

Builds a small REAL .xlsx fixture with openpyxl in this test file (per
task instructions), plus a CSV variant, and drives the module against both.
Negative controls prove absent/malformed input never fabricates a register.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from dv_harness import env_manifest, register_excel_extract as rex


HEADERS = [
    "Register Name", "Offset", "Width", "Access", "Reset Value",
    "Field Name", "Bits", "Field Access", "Field Reset", "Notes",
]


def _write_xlsx(path: Path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    wb.save(path)


def _clean_rows():
    return [
        # register-defining row (also carries an implicit whole-register field concept
        # via a subsequent field row)
        ["CTRL_REG", "0x10", 32, "RW", "0x0", None, None, None, None, "control register"],
        [None, None, None, None, None, "ENABLE", "[0:0]", "RW", "0x0", "enable bit"],
        [None, None, None, None, None, "MODE", "[3:1]", "RW", "0x0", "mode select"],
        ["STATUS_REG", "0x14", 32, "RO", None, None, None, None, None, "status register"],
        [None, None, None, None, None, "DONE_FLAG", "[0]", "RC", "0x0", "read-clears on read"],
        # a field access type the current schema does NOT declare -- RS = read-to-set
        [None, None, None, None, None, "ARM_ON_READ", "[1]", "RS", "0x0", "sets on read"],
        ["IRQ_REG", "0x18", 16, "W1C", "0x0", None, None, None, None, "interrupt clear register"],
    ]


@pytest.fixture()
def clean_xlsx(tmp_path):
    p = tmp_path / "ctrl_regs.xlsx"
    _write_xlsx(p, _clean_rows())
    return p


# ---------------------------------------------------------------------------
# Positive path
# ---------------------------------------------------------------------------

def test_clean_workbook_extracts_all_registers_and_fields(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    assert result.status == rex.STATUS_OK
    assert result.reason is None
    assert result.row_errors == []
    names = [r.register_name for r in result.registers]
    assert names == ["CTRL_REG", "STATUS_REG", "IRQ_REG"]

    ctrl = result.registers[0]
    assert ctrl.offset == 0x10
    assert ctrl.width == 32
    assert ctrl.access_type == "RW"
    assert ctrl.access_in_current_schema is True
    assert ctrl.reset_value == 0x0
    assert [f.name for f in ctrl.fields] == ["ENABLE", "MODE"]
    assert ctrl.fields[0].bit_offset == 0 and ctrl.fields[0].bit_width == 1
    assert ctrl.fields[1].bit_offset == 1 and ctrl.fields[1].bit_width == 3

    status_reg = result.registers[1]
    assert status_reg.access_type == "RO"
    assert status_reg.reset_value is None  # genuinely absent, not fabricated
    done_flag, arm_on_read = status_reg.fields
    assert done_flag.access_type == "RC" and done_flag.access_in_current_schema is True
    assert arm_on_read.access_type == "RS" and arm_on_read.access_in_current_schema is False

    irq = result.registers[2]
    assert irq.width == 16
    assert irq.access_type == "W1C"


def test_schema_widening_candidate_is_reported_for_rs(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    assert "RS" in result.schema_widening_candidates
    assert any("ARM_ON_READ" in where for where in result.schema_widening_candidates["RS"])
    # RS really is absent from the CURRENT schema enum -- this assertion would
    # catch a future schema widening silently making this test meaningless.
    assert "RS" not in result.schema_access_enum_used


def test_absolute_address_computed_when_base_address_supplied(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx, base_address=0x1000)
    ctrl = result.registers[0]
    assert ctrl.absolute_address == 0x1000 + 0x10


def test_absolute_address_none_without_base_address(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    assert result.registers[0].absolute_address is None


def test_csv_variant_parses_identically(tmp_path):
    import csv as csv_mod
    p = tmp_path / "ctrl_regs.csv"
    with p.open("w", newline="", encoding="utf-8") as fh:
        w = csv_mod.writer(fh)
        w.writerow(HEADERS)
        for row in _clean_rows():
            w.writerow(["" if c is None else c for c in row])
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_OK
    assert [r.register_name for r in result.registers] == ["CTRL_REG", "STATUS_REG", "IRQ_REG"]


# ---------------------------------------------------------------------------
# Negative controls -- absent/malformed input must never fabricate a register.
# ---------------------------------------------------------------------------

def test_missing_file_reports_not_available(tmp_path):
    result = rex.extract_register_map(tmp_path / "does_not_exist.xlsx")
    assert result.status == rex.STATUS_NOT_AVAILABLE
    assert "not found" in result.reason
    assert result.registers == []


def test_openpyxl_unavailable_reports_not_available(monkeypatch, clean_xlsx):
    monkeypatch.setattr(rex, "_OPENPYXL_AVAILABLE", False)
    monkeypatch.setattr(rex, "_OPENPYXL_IMPORT_ERROR", "simulated: no module named openpyxl")
    result = rex.extract_register_map(clean_xlsx)
    assert result.status == rex.STATUS_NOT_AVAILABLE
    assert "openpyxl is not installed" in result.reason
    assert "simulated" in result.reason
    assert result.registers == []


def test_legacy_xls_extension_reports_not_available(tmp_path):
    p = tmp_path / "old.xls"
    p.write_bytes(b"not a real xls, just a placeholder")
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_NOT_AVAILABLE
    assert "legacy binary .xls" in result.reason


def test_unsupported_extension_reports_parse_error(tmp_path):
    p = tmp_path / "regs.txt"
    p.write_text("Register Name,Offset\nFOO,0x0\n", encoding="utf-8")
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARSE_ERROR
    assert "unsupported file extension" in result.reason


def test_empty_workbook_reports_parse_error(tmp_path):
    p = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(p)
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARSE_ERROR
    assert "empty" in result.reason


def test_missing_required_headers_reports_parse_error(tmp_path):
    p = tmp_path / "bad_headers.xlsx"
    _write_xlsx(p, [])
    wb = openpyxl.load_workbook(p)
    ws = wb.active
    ws.delete_rows(1, ws.max_row)
    ws.append(["Some Column", "Another Column"])
    ws.append(["x", "y"])
    wb.save(p)
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARSE_ERROR
    assert "missing required column" in result.reason


def test_all_rows_malformed_reports_parse_error_with_no_registers(tmp_path):
    p = tmp_path / "all_bad.xlsx"
    _write_xlsx(p, [
        ["BAD_REG", "not-hex-and-not-decimal", 32, "RW", "0x0", None, None, None, None, None],
    ])
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARSE_ERROR
    assert result.registers == []
    assert len(result.row_errors) == 1
    assert "invalid offset" in result.row_errors[0]["error"]


def test_bad_offset_row_excluded_but_good_rows_still_extracted_partial(tmp_path):
    p = tmp_path / "mixed.xlsx"
    _write_xlsx(p, [
        ["GOOD_REG", "0x0", 32, "RW", "0x0", None, None, None, None, None],
        ["BAD_REG", "not-a-number", 32, "RW", "0x0", None, None, None, None, None],
    ])
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARTIAL
    assert [r.register_name for r in result.registers] == ["GOOD_REG"]
    assert len(result.row_errors) == 1
    assert result.row_errors[0]["register"] == "BAD_REG"


def test_bad_bit_range_excludes_field_but_keeps_register(tmp_path):
    p = tmp_path / "bad_bits.xlsx"
    _write_xlsx(p, [
        ["REG1", "0x0", 32, "RW", "0x0", None, None, None, None, None],
        [None, None, None, None, None, "GARBLED_FIELD", "not-a-bit-range", "RW", "0x0", None],
    ])
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARTIAL
    assert len(result.registers) == 1
    assert result.registers[0].fields == []
    assert any("bit range" in e["error"] for e in result.row_errors)


def test_field_with_no_preceding_register_is_a_row_error(tmp_path):
    p = tmp_path / "orphan_field.xlsx"
    _write_xlsx(p, [
        [None, None, None, None, None, "ORPHAN", "[0]", "RW", "0x0", None],
    ])
    result = rex.extract_register_map(p)
    assert result.registers == []
    assert any("no preceding register" in e["error"] for e in result.row_errors)


def test_unrecognized_access_text_is_reported_not_silently_dropped(tmp_path):
    p = tmp_path / "garbage_access.xlsx"
    _write_xlsx(p, [
        ["REG1", "0x0", 32, "this is not an access mnemonic at all", "0x0", None, None, None, None, None],
    ])
    result = rex.extract_register_map(p)
    assert result.status == rex.STATUS_PARTIAL
    assert result.registers[0].access_type is None
    assert any("unrecognized access-type" in e["error"] for e in result.row_errors)


def test_sheet_name_not_found_reports_parse_error(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx, sheet_name="NoSuchSheet")
    assert result.status == rex.STATUS_PARSE_ERROR
    assert "NoSuchSheet" in result.reason


# ---------------------------------------------------------------------------
# Access-type normalization
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("raw,expected", [
    ("RW", "RW"), ("R/W", "RW"), ("Read/Write", "RW"), ("read write", "RW"),
    ("RO", "RO"), ("R", "RO"), ("Read Only", "RO"),
    ("WO", "WO"), ("Write Only", "WO"),
    ("W1C", "W1C"), ("Write 1 to Clear", "W1C"), ("w1tc", "W1C"),
    ("RW1C", "RW1C"),
    ("RC", "RC"),
    ("WC", "WC"),
    ("W1S", "W1S"),
    ("RS", "RS"),
])
def test_normalize_access_type_aliases(raw, expected):
    token, err = rex.normalize_access_type(raw)
    assert err is None
    assert token == expected


def test_normalize_access_type_blank_is_none():
    assert rex.normalize_access_type(None) == (None, None)
    assert rex.normalize_access_type("") == (None, None)
    assert rex.normalize_access_type("N/A") == (None, None)


def test_normalize_access_type_garbage_is_error():
    token, err = rex.normalize_access_type("please read the datasheet section 4.2 for details")
    assert token is None
    assert err is not None


# ---------------------------------------------------------------------------
# schema_access_kind_enum() reads the real schema file, not a hardcoded list.
# ---------------------------------------------------------------------------

def test_schema_access_kind_enum_matches_real_schema_file():
    enum = rex.schema_access_kind_enum()
    assert set(enum) == {"RW", "RO", "WO", "W1C", "RW1C", "RC", "WC", "W1S"}


# ---------------------------------------------------------------------------
# Bridge to register_map.schema.json / env_manifest.validate_register_map()
# ---------------------------------------------------------------------------

def test_to_register_map_document_validates_against_real_schema(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    bridged = rex.to_register_map_document(result)
    # validate_register_map() would have raised on failure -- reaching here
    # already proves it passed. Re-assert explicitly for a clear failure message.
    env_manifest.validate_register_map(bridged["document"])


def test_to_register_map_document_excludes_schema_widening_candidates(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    bridged = rex.to_register_map_document(result)
    doc = bridged["document"]
    all_field_names = {
        f["name"]
        for block in doc["blocks"]
        for reg in block["registers"]
        for f in reg["fields"]
    }
    # ARM_ON_READ's access type (RS) is not in the current schema enum --
    # it must be excluded from the schema-conformant document...
    assert "ARM_ON_READ" not in all_field_names
    # ...but DONE_FLAG (RC, a real schema value) must survive.
    assert "DONE_FLAG" in all_field_names
    excluded_fields = {e.get("field") for e in bridged["excluded"] if "field" in e}
    assert "ARM_ON_READ" in excluded_fields


def test_to_register_map_document_excludes_register_with_bad_width(clean_xlsx):
    result = rex.extract_register_map(clean_xlsx)
    # Corrupt one register's width to a value the schema does not declare.
    result.registers[0].width = 24
    bridged = rex.to_register_map_document(result)
    doc = bridged["document"]
    all_reg_names = {reg["name"] for block in doc["blocks"] for reg in block["registers"]}
    assert "CTRL_REG" not in all_reg_names
    assert any(e.get("register") == "CTRL_REG" and "width" in e["reason"] for e in bridged["excluded"])


# ---------------------------------------------------------------------------
# CLI front door (ad hoc `python -m`, per file-safety scope -- no cli.py verb).
# ---------------------------------------------------------------------------

def test_cli_exit_code_zero_on_clean_workbook(clean_xlsx):
    proc = subprocess.run(
        [sys.executable, "-m", "dv_harness.register_excel_extract", str(clean_xlsx), "--json"],
        cwd=str(Path(__file__).resolve().parent.parent),
        capture_output=True, text=True,
    )
    assert proc.returncode == 0, proc.stdout + proc.stderr
    assert '"status": "OK"' in proc.stdout


def test_cli_exit_code_two_on_missing_file(tmp_path):
    proc = subprocess.run(
        [sys.executable, "-m", "dv_harness.register_excel_extract", str(tmp_path / "nope.xlsx")],
        cwd=str(Path(__file__).resolve().parent.parent),
        capture_output=True, text=True,
    )
    assert proc.returncode == 2
    assert "NOT_AVAILABLE" in proc.stdout
